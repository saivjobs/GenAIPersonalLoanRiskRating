from pptx import Presentation
from pptx.util import Inches, Pt
from pathlib import Path

root = Path(__file__).resolve().parents[1]
pptx_out = root / 'ARCHITECTURE_SLIDE_OPTIMIZED.pptx'
img = root / 'ARCHITECTURE.png'

prs = Presentation()
slide_layout = prs.slide_layouts[6]  # blank
slide = prs.slides.add_slide(slide_layout)

# Title
title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(9), Inches(1))
title_tf = title_box.text_frame
title_tf.clear()
p = title_tf.paragraphs[0]
p.text = "Gen AI Personal Loan Risk Rating — Architecture"
p.font.size = Pt(28)
p.font.bold = True

# Left bullets
left = Inches(0.4)
top = Inches(1.2)
body = slide.shapes.add_textbox(left, top, Inches(5), Inches(4)).text_frame
body.word_wrap = True
body.clear()
bullets = [
    "Streamlit UI (app.py) → Orchestration (main.py / workflow.py)",
    "Scoring: GradientBoostingClassifier + decile mapping",
    "Document extraction: document_intelligence.py",
    "RAG: Gemini embeddings → Chroma vector store → retriever",
    "LLM: Gemini for reasoning; risk_scoring.py builds prompt",
    "Persistence: SQLite audit log + model_artifacts/",
]
for i, b in enumerate(bullets):
    p = body.add_paragraph()
    p.text = b
    p.level = 0
    p.font.size = Pt(13)

# Right image
if img.exists():
    left = Inches(5.8)
    top = Inches(1.2)
    width = Inches(3.2)
    slide.shapes.add_picture(str(img), left, top, width=width)

# Speaker notes (concise)
notes = (
    "Architecture overview: \n"
    "- Uses a trained PD model (scikit-learn) to compute PD, mapped to deciles.\n"
    "- RAG layer grounds Gemini LLM reasoning in policy docs via Chroma.\n"
    "- Audit log (SQLite) stores full context; human overrides supported.\n"
    "Demo tip: show PD, decile, policy sources cited, audit log, and fairness screen."
)
notes_slide = slide.notes_slide
notes_tf = notes_slide.notes_text_frame
notes_tf.clear()
notes_tf.text = notes

prs.save(pptx_out)
print(f"Wrote optimized PPTX: {pptx_out}")
