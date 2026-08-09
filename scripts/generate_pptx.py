from pptx import Presentation
from pptx.util import Inches, Pt
from pathlib import Path

root = Path(__file__).resolve().parents[1]
pptx_path = root / "ARCHITECTURE_SLIDE.pptx"
md_path = root / "ARCHITECTURE.mmd"

prs = Presentation()
# Single blank slide layout
slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(slide_layout)

# Title
title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
title_tf = title_box.text_frame
p = title_tf.paragraphs[0]
p.text = "Gen AI Personal Loan Risk Rating — Architecture"
p.font.size = Pt(28)

# Bullet points with key components
left = Inches(0.5)
top = Inches(1.3)
width = Inches(4.5)
height = Inches(4)
body = slide.shapes.add_textbox(left, top, width, height).text_frame
body.word_wrap = True
for point in [
    "Streamlit UI (app.py) connects to orchestration (main.py / workflow.py)",
    "Scoring: trained GradientBoostingClassifier (pd_model.py) + decile mapping",
    "Document extraction: document_intelligence.py produces structured facts",
    "RAG: Gemini embeddings -> Chroma vector store (rag_index) -> retriever",
    "LLM: Gemini via gemini_client.py; risk_scoring.py builds prompt and parses response",
    "Persistence: SQLite audit log (db.py) and model_artifacts/ for the model",
]:
    p = body.add_paragraph()
    p.text = point
    p.level = 0
    p.font.size = Pt(14)

# Include Mermaid source as a small code box on right
mermaid_text = md_path.read_text(encoding='utf-8') if md_path.exists() else "(mermaid diagram not found)"
right = Inches(5.2)
code_box = slide.shapes.add_textbox(right, top, Inches(4.3), Inches(4))
code_tf = code_box.text_frame
code_tf.word_wrap = True
code_tf.margin_left = Pt(6)
code_tf.margin_top = Pt(6)
code_tf.margin_right = Pt(6)
code_tf.margin_bottom = Pt(6)
code_tf.paragraphs[0].font.size = Pt(10)
# Add mermaid text
for line in mermaid_text.splitlines():
    p = code_tf.add_paragraph()
    p.text = line
    p.level = 0
    p.font.size = Pt(8)

prs.save(pptx_path)
print(f"Wrote PPTX: {pptx_path}")
