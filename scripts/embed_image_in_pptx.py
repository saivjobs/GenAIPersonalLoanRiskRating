from pathlib import Path
from pptx import Presentation

root = Path(__file__).resolve().parents[1]
pptx_in = root / 'ARCHITECTURE_SLIDE.pptx'
img = root / 'ARCHITECTURE.png'
pptx_out = root / 'ARCHITECTURE_SLIDE_WITH_IMAGE.pptx'

if not pptx_in.exists():
    raise SystemExit('PPTX input not found')
if not img.exists():
    raise SystemExit('Image not found')

prs = Presentation(str(pptx_in))
slide = prs.slides[0]
# Add image on the right, same position as earlier code box
left = prs.slide_width * 0.55
top = prs.slide_height * 0.15
width = prs.slide_width * 0.4
slide.shapes.add_picture(str(img), left, top, width=width)
prs.save(str(pptx_out))
print('Wrote', pptx_out)
